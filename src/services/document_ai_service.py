"""
Document AI Service with RAG (Retrieval-Augmented Generation)
REAL implementation with document processing, embeddings, and GPT-4 integration
"""
import os
import io
import hashlib
import mimetypes
from pathlib import Path
from typing import List, Dict, Any, Optional, Tuple
import logging
from datetime import datetime
import json

import openai
import tiktoken
from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain.embeddings import OpenAIEmbeddings
from langchain.vectorstores import Chroma
import PyPDF2
import docx
import openpyxl
from PIL import Image
import pytesseract
import aiofiles
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, and_

from src.db.models.device_document import DeviceDocument
from src.db.models.document_chunk import DocumentChunk
from src.db.models.chat import ChatConversation, ChatMessage, ChatCitation
from src.core.config import settings
from src.services.s3_service import S3Service

logger = logging.getLogger(__name__)

# Initialize OpenAI
openai.api_key = settings.OPENAI_API_KEY


class DocumentAIService:
    """
    REAL document processing and AI service.
    Handles document upload, text extraction, chunking, embeddings, and RAG.
    """
    
    def __init__(self):
        self.s3_service = S3Service()
        self.embeddings = OpenAIEmbeddings(openai_api_key=settings.OPENAI_API_KEY)
        self.encoding = tiktoken.encoding_for_model("gpt-4")
        
        # Text splitter for chunking
        self.text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=1000,
            chunk_overlap=200,
            length_function=self.count_tokens,
            separators=["\n\n", "\n", ". ", " ", ""]
        )
        
        # Supported file types
        self.supported_types = {
            'application/pdf': self.extract_pdf,
            'application/vnd.openxmlformats-officedocument.wordprocessingml.document': self.extract_docx,
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet': self.extract_xlsx,
            'text/plain': self.extract_text,
            'image/jpeg': self.extract_image_ocr,
            'image/png': self.extract_image_ocr,
        }
    
    def count_tokens(self, text: str) -> int:
        """Count tokens in text using tiktoken."""
        return len(self.encoding.encode(text))
    
    async def process_document(
        self,
        db: AsyncSession,
        file_path: str,
        file_name: str,
        device_id: Optional[int] = None,
        organization_id: int = None,
        user_id: int = None,
        document_type: str = "manual"
    ) -> DeviceDocument:
        """
        Process uploaded document - extract text, create chunks, generate embeddings.
        """
        logger.info(f"Processing document: {file_name}")
        
        # Get file info
        file_size = os.path.getsize(file_path)
        mime_type, _ = mimetypes.guess_type(file_name)
        
        # Calculate file hash
        file_hash = await self._calculate_file_hash(file_path)
        
        # Check if document already exists
        existing = await db.execute(
            select(DeviceDocument).where(
                DeviceDocument.file_hash == file_hash
            )
        )
        if existing.scalar_one_or_none():
            logger.info(f"Document already exists with hash: {file_hash}")
            return existing.scalar_one()
        
        # Upload to S3
        s3_key = f"documents/{organization_id}/{file_hash}/{file_name}"
        s3_url = await self.s3_service.upload_file(file_path, s3_key)
        
        # Extract text based on file type
        text_content = await self._extract_text(file_path, mime_type)
        
        # Create document record
        document = DeviceDocument(
            device_id=device_id,
            organization_id=organization_id,
            document_type=document_type,
            title=file_name,
            file_url=s3_url,
            file_key=s3_key,
            file_size_bytes=file_size,
            file_hash=file_hash,
            mime_type=mime_type,
            page_count=self._count_pages(file_path, mime_type),
            created_by=user_id
        )
        db.add(document)
        await db.flush()  # Get document ID
        
        # Create chunks and embeddings
        chunks = await self._create_chunks(text_content)
        chunk_embeddings = await self._generate_embeddings([chunk['text'] for chunk in chunks])
        
        # Store chunks in database
        for i, (chunk, embedding) in enumerate(zip(chunks, chunk_embeddings)):
            chunk_record = DocumentChunk(
                document_id=document.id,
                chunk_index=i,
                chunk_text=chunk['text'],
                page_number=chunk.get('page', 1),
                section_heading=chunk.get('section'),
                embedding_model='text-embedding-ada-002',
                embedding=embedding,
                metadata=json.dumps(chunk.get('metadata', {}))
            )
            db.add(chunk_record)
        
        await db.commit()
        logger.info(f"Document processed: {document.id} with {len(chunks)} chunks")
        
        return document
    
    async def _calculate_file_hash(self, file_path: str) -> str:
        """Calculate SHA-256 hash of file."""
        sha256_hash = hashlib.sha256()
        async with aiofiles.open(file_path, 'rb') as f:
            while chunk := await f.read(8192):
                sha256_hash.update(chunk)
        return sha256_hash.hexdigest()
    
    async def _extract_text(self, file_path: str, mime_type: str) -> str:
        """Extract text from document based on mime type."""
        extractor = self.supported_types.get(mime_type)
        if not extractor:
            raise ValueError(f"Unsupported file type: {mime_type}")
        
        return await extractor(file_path)
    
    async def extract_pdf(self, file_path: str) -> str:
        """Extract text from PDF."""
        text_parts = []
        
        with open(file_path, 'rb') as file:
            pdf_reader = PyPDF2.PdfReader(file)
            
            for page_num, page in enumerate(pdf_reader.pages):
                try:
                    text = page.extract_text()
                    if text.strip():
                        text_parts.append(f"[Page {page_num + 1}]\n{text}")
                except Exception as e:
                    logger.error(f"Error extracting page {page_num}: {e}")
        
        return "\n\n".join(text_parts)
    
    async def extract_docx(self, file_path: str) -> str:
        """Extract text from Word document."""
        doc = docx.Document(file_path)
        text_parts = []
        
        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text)
        
        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells)
                if row_text.strip():
                    text_parts.append(row_text)
        
        return "\n\n".join(text_parts)
    
    async def extract_xlsx(self, file_path: str) -> str:
        """Extract text from Excel file."""
        wb = openpyxl.load_workbook(file_path, read_only=True)
        text_parts = []
        
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            text_parts.append(f"[Sheet: {sheet_name}]")
            
            for row in sheet.iter_rows(values_only=True):
                row_text = " | ".join(str(cell) for cell in row if cell)
                if row_text.strip():
                    text_parts.append(row_text)
        
        return "\n\n".join(text_parts)
    
    async def extract_text(self, file_path: str) -> str:
        """Extract text from plain text file."""
        async with aiofiles.open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            return await f.read()
    
    async def extract_image_ocr(self, file_path: str) -> str:
        """Extract text from image using OCR."""
        try:
            image = Image.open(file_path)
            text = pytesseract.image_to_string(image)
            return text
        except Exception as e:
            logger.error(f"OCR error: {e}")
            return ""
    
    def _count_pages(self, file_path: str, mime_type: str) -> int:
        """Count pages in document."""
        if mime_type == 'application/pdf':
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                return len(pdf_reader.pages)
        return 1
    
    async def _create_chunks(self, text: str) -> List[Dict[str, Any]]:
        """Create text chunks with metadata."""
        # Split text into chunks
        chunks = self.text_splitter.split_text(text)
        
        # Add metadata to each chunk
        chunk_data = []
        for i, chunk_text in enumerate(chunks):
            # Extract section heading if possible
            lines = chunk_text.split('\n')
            section = None
            if lines and len(lines[0]) < 100 and lines[0].strip():
                section = lines[0].strip()
            
            chunk_data.append({
                'text': chunk_text,
                'index': i,
                'section': section,
                'metadata': {
                    'chunk_size': len(chunk_text),
                    'token_count': self.count_tokens(chunk_text)
                }
            })
        
        return chunk_data
    
    async def _generate_embeddings(self, texts: List[str]) -> List[List[float]]:
        """Generate embeddings for text chunks."""
        embeddings = []
        
        # Process in batches to avoid rate limits
        batch_size = 20
        for i in range(0, len(texts), batch_size):
            batch = texts[i:i + batch_size]
            batch_embeddings = await self.embeddings.aembed_documents(batch)
            embeddings.extend(batch_embeddings)
        
        return embeddings
    
    async def query_documents(
        self,
        db: AsyncSession,
        query: str,
        device_id: Optional[int] = None,
        organization_id: Optional[int] = None,
        top_k: int = 5
    ) -> List[Tuple[DocumentChunk, float]]:
        """
        Query documents using semantic search.
        Returns relevant chunks with similarity scores.
        """
        # Generate query embedding
        query_embedding = await self.embeddings.aembed_query(query)
        
        # Build base query
        stmt = select(DocumentChunk).join(DeviceDocument)
        
        # Add filters
        conditions = []
        if device_id:
            conditions.append(DeviceDocument.device_id == device_id)
        if organization_id:
            conditions.append(DeviceDocument.organization_id == organization_id)
        
        if conditions:
            stmt = stmt.where(and_(*conditions))
        
        # Execute query
        result = await db.execute(stmt)
        chunks = result.scalars().all()
        
        # Calculate similarities
        chunk_scores = []
        for chunk in chunks:
            # Cosine similarity
            similarity = self._cosine_similarity(query_embedding, chunk.embedding)
            chunk_scores.append((chunk, similarity))
        
        # Sort by similarity and return top k
        chunk_scores.sort(key=lambda x: x[1], reverse=True)
        return chunk_scores[:top_k]
    
    def _cosine_similarity(self, vec1: List[float], vec2: List[float]) -> float:
        """Calculate cosine similarity between two vectors."""
        import numpy as np
        vec1 = np.array(vec1)
        vec2 = np.array(vec2)
        return np.dot(vec1, vec2) / (np.linalg.norm(vec1) * np.linalg.norm(vec2))
    
    async def generate_ai_response(
        self,
        db: AsyncSession,
        conversation_id: int,
        user_message: str,
        device_context: Optional[Dict[str, Any]] = None
    ) -> ChatMessage:
        """
        Generate AI response using GPT-4 with RAG.
        """
        conversation = await db.get(ChatConversation, conversation_id)
        if not conversation:
            raise ValueError("Conversation not found")
        
        # Query relevant documents
        relevant_chunks = []
        if conversation.device_id:
            relevant_chunks = await self.query_documents(
                db,
                user_message,
                device_id=conversation.device_id,
                top_k=5
            )
        
        # Build context from relevant documents
        context_parts = []
        citations = []
        
        for i, (chunk, score) in enumerate(relevant_chunks):
            if score > 0.7:  # Relevance threshold
                context_parts.append(f"[Source {i+1}]: {chunk.chunk_text}")
                
                # Get document info for citation
                document = await db.get(DeviceDocument, chunk.document_id)
                citations.append({
                    'chunk_id': chunk.id,
                    'document_id': document.id,
                    'document_title': document.title,
                    'page_number': chunk.page_number,
                    'relevance_score': score
                })
        
        # Add device context if provided
        if device_context:
            context_parts.insert(0, f"Device Information:\n{json.dumps(device_context, indent=2)}")
        
        # Build messages for GPT-4
        messages = [
            {
                "role": "system",
                "content": """You are a medical device expert assistant. 
                Provide accurate, helpful information based on the provided context.
                Always cite your sources using [Source N] references.
                If information is not in the context, say so clearly.
                Never provide medical diagnosis or treatment advice."""
            }
        ]
        
        # Add conversation history
        recent_messages = await db.execute(
            select(ChatMessage)
            .where(ChatMessage.conversation_id == conversation_id)
            .order_by(ChatMessage.created_at.desc())
            .limit(10)
        )
        
        for msg in reversed(recent_messages.scalars().all()):
            messages.append({
                "role": msg.role,
                "content": msg.content
            })
        
        # Add context and user message
        if context_parts:
            context_content = "\n\n".join(context_parts)
            messages.append({
                "role": "system",
                "content": f"Relevant context:\n{context_content}"
            })
        
        messages.append({
            "role": "user",
            "content": user_message
        })
        
        # Call GPT-4
        try:
            response = await openai.ChatCompletion.acreate(
                model="gpt-4",
                messages=messages,
                max_tokens=2000,
                temperature=0.3,
                stream=False
            )
            
            ai_content = response.choices[0].message.content
            tokens_used = response.usage.total_tokens
            
            # Calculate confidence score based on context relevance
            confidence_score = 0.0
            if relevant_chunks:
                confidence_score = sum(score for _, score in relevant_chunks[:3]) / min(3, len(relevant_chunks))
            
            # Create message record
            message = ChatMessage(
                conversation_id=conversation_id,
                role="assistant",
                content=ai_content,
                tokens_used=tokens_used,
                model_used="gpt-4",
                has_citations=len(citations) > 0,
                confidence_score=confidence_score
            )
            db.add(message)
            await db.flush()
            
            # Create citations
            for citation_data in citations:
                citation = ChatCitation(
                    message_id=message.id,
                    source_type="document",
                    source_id=str(citation_data['document_id']),
                    source_title=citation_data['document_title'],
                    page_number=citation_data['page_number'],
                    relevance_score=citation_data['relevance_score']
                )
                db.add(citation)
            
            # Update conversation
            conversation.total_messages += 2  # User + Assistant
            conversation.total_tokens_used += tokens_used
            conversation.last_message_at = datetime.utcnow()
            
            await db.commit()
            
            return message
            
        except Exception as e:
            logger.error(f"GPT-4 API error: {e}")
            raise
    
    async def export_conversation(
        self,
        db: AsyncSession,
        conversation_id: int,
        format: str = "pdf"
    ) -> str:
        """
        Export conversation with citations.
        """
        # Get conversation with all messages and citations
        conversation = await db.get(ChatConversation, conversation_id)
        
        messages = await db.execute(
            select(ChatMessage)
            .where(ChatMessage.conversation_id == conversation_id)
            .order_by(ChatMessage.created_at)
        )
        
        # Build export content
        content = f"# {conversation.title}\n\n"
        content += f"Created: {conversation.created_at}\n\n"
        
        for message in messages.scalars().all():
            content += f"**{message.role.title()}**: {message.content}\n\n"
            
            # Add citations
            citations = await db.execute(
                select(ChatCitation)
                .where(ChatCitation.message_id == message.id)
            )
            
            citation_list = citations.scalars().all()
            if citation_list:
                content += "Sources:\n"
                for citation in citation_list:
                    content += f"- {citation.source_title}"
                    if citation.page_number:
                        content += f" (Page {citation.page_number})"
                    content += "\n"
                content += "\n"
        
        # Generate file based on format
        if format == "pdf":
            # TODO: Implement PDF generation
            pass
        else:
            # Return markdown for now
            return content
